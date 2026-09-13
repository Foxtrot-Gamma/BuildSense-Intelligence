import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_baseline_excel(finalized_schedule):
    """
    Generate a professionally formatted Excel workbook
    containing the finalized project baseline.
    """

    required_columns = [
        "ID",
        "WBS",
        "Activity",
        "Duration",
        "Duration Source",
        "Relationship",
        "Predecessor",
        "Planned Start",
        "Planned Finish",
        "Total Float",
        "Critical",
    ]

    # Validate input
    if finalized_schedule is None:
        return None

    if finalized_schedule.empty:
        return None

    missing_columns = [
        col for col in required_columns
        if col not in finalized_schedule.columns
    ]

    if missing_columns:
        return None

    # Work on a copy
    schedule = finalized_schedule[required_columns].copy()

    # Convert dates
    schedule["Planned Start"] = pd.to_datetime(
        schedule["Planned Start"],
        errors="coerce"
    )

    schedule["Planned Finish"] = pd.to_datetime(
        schedule["Planned Finish"],
        errors="coerce"
    )

    # Calculate project summary
    valid_starts = schedule["Planned Start"].dropna()
    valid_finishes = schedule["Planned Finish"].dropna()

    if valid_starts.empty or valid_finishes.empty:
        return None

    project_start = valid_starts.min()
    project_finish = valid_finishes.max()

    project_duration = (
        project_finish - project_start
    ).days + 1

    activity_count = len(schedule)

    numeric_float = pd.to_numeric(
        schedule["Total Float"],
        errors="coerce"
    )

    maximum_float = (
        numeric_float.max()
        if not numeric_float.dropna().empty
        else 0
    )

    critical_values = schedule["Critical"].astype(str).str.lower()

    critical_count = critical_values.isin(
        ["true", "1", "yes"]
    ).sum()

    # Create workbook
    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # =========================================================
    # SHEET 1 — BASELINE SUMMARY
    # =========================================================

    summary = workbook.create_sheet("Baseline Summary")

    # Title
    summary["A1"] = "PROJECT BASELINE — SCHEDULE SUMMARY"
    summary["A1"].font = Font(
        bold=True,
        size=16
    )

    summary.merge_cells("A1:B1")

    # Project name
    try:
        import streamlit as st

        project_name = st.session_state.get(
            "project_name",
            "Construction Project"
        )
    except Exception:
        project_name = "Construction Project"

    summary["A3"] = "Project"
    summary["B3"] = project_name

    summary["A4"] = "Project Start"
    summary["B4"] = project_start

    summary["A5"] = "Project Finish"
    summary["B5"] = project_finish

    summary["A6"] = "Project Duration"
    summary["B6"] = project_duration

    summary["A7"] = "Total Activities"
    summary["B7"] = activity_count

    summary["A8"] = "Critical Activities"
    summary["B8"] = critical_count

    summary["A9"] = "Maximum Activity Float"
    summary["B9"] = maximum_float

    # Summary formatting
    label_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2"
    )

    for row in range(3, 10):
        summary[f"A{row}"].font = Font(bold=True)
        summary[f"A{row}"].fill = label_fill

    summary["B4"].number_format = "dd-mmm-yyyy"
    summary["B5"].number_format = "dd-mmm-yyyy"

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 30

    # =========================================================
    # SHEET 2 — BASELINE SCHEDULE
    # =========================================================

    schedule_sheet = workbook.create_sheet("Baseline Schedule")

    # Write headers
    for col_index, column_name in enumerate(
        required_columns,
        start=1
    ):
        cell = schedule_sheet.cell(
            row=1,
            column=col_index,
            value=column_name
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Write schedule data
    for row_index, row in enumerate(
        schedule.itertuples(index=False),
        start=2
    ):

        for col_index, value in enumerate(
            row,
            start=1
        ):

            cell = schedule_sheet.cell(
                row=row_index,
                column=col_index,
                value=value
            )

            cell.alignment = Alignment(
                vertical="center"
            )

    # Date formatting
    start_col = required_columns.index(
        "Planned Start"
    ) + 1

    finish_col = required_columns.index(
        "Planned Finish"
    ) + 1

    for row in range(
        2,
        schedule_sheet.max_row + 1
    ):
        schedule_sheet.cell(
            row=row,
            column=start_col
        ).number_format = "dd-mmm-yyyy"

        schedule_sheet.cell(
            row=row,
            column=finish_col
        ).number_format = "dd-mmm-yyyy"

    # Freeze header
    schedule_sheet.freeze_panes = "A2"

    # Enable filters
    schedule_sheet.auto_filter.ref = (
        schedule_sheet.dimensions
    )

    # Column widths
    widths = {
        "ID": 10,
        "WBS": 15,
        "Activity": 40,
        "Duration": 12,
        "Duration Source": 20,
        "Relationship": 15,
        "Predecessor": 18,
        "Planned Start": 17,
        "Planned Finish": 17,
        "Total Float": 14,
        "Critical": 12,
    }

    for index, column_name in enumerate(
        required_columns,
        start=1
    ):
        schedule_sheet.column_dimensions[
            get_column_letter(index)
        ].width = widths.get(
            column_name,
            15
        )

    # =========================================================
    # CRITICAL ACTIVITY HIGHLIGHTING
    # =========================================================

    critical_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )

    critical_font = Font(
        bold=True
    )

    critical_column = required_columns.index(
        "Critical"
    ) + 1

    for row in range(
        2,
        schedule_sheet.max_row + 1
    ):

        critical_value = str(
            schedule_sheet.cell(
                row=row,
                column=critical_column
            ).value
        ).lower()

        if critical_value in [
            "true",
            "1",
            "yes"
        ]:

            for col in range(
                1,
                schedule_sheet.max_column + 1
            ):
                schedule_sheet.cell(
                    row=row,
                    column=col
                ).fill = critical_fill

                schedule_sheet.cell(
                    row=row,
                    column=col
                ).font = critical_font

    # =========================================================
    # BORDER / ALIGNMENT
    # =========================================================

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9D9D9"
        )
    )

    for row in schedule_sheet.iter_rows():
        for cell in row:
            cell.border = thin_border

    schedule_sheet.row_dimensions[1].height = 24

    # =========================================================
    # EXPORT TO MEMORY
    # =========================================================

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    return output